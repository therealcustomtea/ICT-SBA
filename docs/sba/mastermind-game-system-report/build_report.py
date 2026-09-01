#!/usr/bin/env python3
# ruff: noqa: E501
"""Build the HKDSE ICT SBA report from its inspectable Markdown source."""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "report_source.md"
ASSETS = ROOT / "assets"
ICON = ROOT.parents[2] / "apps" / "web" / "public" / "icon-512.png"

NAVY = "17324D"
TEAL = "2B6F77"
CORAL = "D46B52"
SAND = "F5EFE6"
PALE_BLUE = "EAF2F4"
INK = "1F2933"
MUTED = "647383"
WHITE = "FFFFFF"

CANDIDATE_NAME = "CHAN, YU SHING"
CANDIDATE_NUMBER = "Student ID s202101127"
SCHOOL_NAME = "CCC Ming Kei College"
GRID = "C9D6DB"

FIGURE_RE = re.compile(r"^\[\[FIGURE:([^|]+)\|(.+)\]\]$")
HEADING_RE = re.compile(r"^(#{1,3})\s+(.+)$")
TABLE_SEPARATOR_RE = re.compile(r"^\|(?:\s*:?-+:?\s*\|)+$")
INLINE_RE = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|\*[^*]+\*|\[[^]]+\]\([^)]+\))")


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(
    cell, top: int = 80, start: int = 100, bottom: int = 80, end: int = 100
) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def set_table_fixed(table, widths: list[float]) -> None:
    tbl_pr = table._tbl.tblPr
    table.autofit = False
    width_twips = [round(width * 1440) for width in widths]
    total_twips = sum(width_twips)
    tbl_width = tbl_pr.find(qn("w:tblW"))
    if tbl_width is None:
        tbl_width = OxmlElement("w:tblW")
        tbl_pr.append(tbl_width)
    tbl_width.set(qn("w:w"), str(total_twips))
    tbl_width.set(qn("w:type"), "dxa")
    tbl_indent = tbl_pr.find(qn("w:tblInd"))
    if tbl_indent is None:
        tbl_indent = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_indent)
    tbl_indent.set(qn("w:w"), "100")
    tbl_indent.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for twips in width_twips:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(twips))
        grid.append(column)
    for row in table.rows:
        for index, (width, twips) in enumerate(zip(widths, width_twips, strict=True)):
            row.cells[index].width = Inches(width)
            tc_pr = row.cells[index]._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(twips))
            tc_w.set(qn("w:type"), "dxa")


def set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)
    set_repeat_table_header(row)


def create_numbering_instance(document: Document, base_num_id: int = 5) -> int:
    """Create a fresh ordered-list instance so each Markdown block starts at one."""
    root = document.part.numbering_part.element
    existing = [int(node.get(qn("w:numId"))) for node in root.findall(qn("w:num"))]
    new_id = max(existing, default=0) + 1
    abstract_id = "7"
    for node in root.findall(qn("w:num")):
        if node.get(qn("w:numId")) == str(base_num_id):
            abstract = node.find(qn("w:abstractNumId"))
            if abstract is not None:
                abstract_id = abstract.get(qn("w:val"))
            break
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(new_id))
    abstract = OxmlElement("w:abstractNumId")
    abstract.set(qn("w:val"), abstract_id)
    num.append(abstract)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride")
    start.set(qn("w:val"), "1")
    override.append(start)
    num.append(override)
    root.append(num)
    return new_id


def apply_numbering(paragraph, num_id: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    level = OxmlElement("w:ilvl")
    level.set(qn("w:val"), "0")
    number = OxmlElement("w:numId")
    number.set(qn("w:val"), str(num_id))
    num_pr.extend((level, number))


def add_field(run, instruction: str) -> None:
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, text, separate, end))


def add_bookmark(paragraph, name: str, bookmark_id: int) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def add_internal_link(paragraph, text: str, anchor: str, color: str = INK) -> None:
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), anchor)
    hyperlink.set(qn("w:history"), "1")
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    colour = OxmlElement("w:color")
    colour.set(qn("w:val"), color)
    props.append(colour)
    run.append(props)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_external_link(paragraph, text: str, url: str) -> None:
    relationship = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    colour = OxmlElement("w:color")
    colour.set(qn("w:val"), TEAL)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.extend((colour, underline))
    run.append(props)
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_inline(paragraph, text: str) -> None:
    cursor = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > cursor:
            paragraph.add_run(text[cursor : match.start()])
        token = match.group(0)
        if token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Menlo"
            run.font.size = Pt(8.7)
            run.font.color.rgb = RGBColor.from_string(NAVY)
            run._r.get_or_add_rPr().append(OxmlElement("w:noProof"))
        elif token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("*"):
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        else:
            link_match = re.match(r"\[([^]]+)\]\(([^)]+)\)", token)
            if link_match:
                add_external_link(paragraph, link_match.group(1), link_match.group(2))
        cursor = match.end()
    if cursor < len(text):
        paragraph.add_run(text[cursor:])


def set_run_fonts(run, name: str = "Arial") -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def configure_styles(document: Document) -> None:
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.2)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.paragraph_format.space_after = Pt(5.5)
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for style_name, size, colour, before, after in (
        ("Title", 29, NAVY, 0, 12),
        ("Subtitle", 12, MUTED, 4, 10),
        ("Heading 1", 18, NAVY, 0, 9),
        ("Heading 2", 13.5, TEAL, 10, 5),
        ("Heading 3", 11.5, CORAL, 8, 4),
    ):
        style = styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(colour)
        style.font.bold = True
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    styles["Heading 1"].paragraph_format.page_break_before = True
    styles["Heading 1"].paragraph_format.border_bottom = True

    for name, left, hanging in (("List Bullet", 0.28, 0.18), ("List Number", 0.28, 0.18)):
        style = styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(10)
        style.paragraph_format.left_indent = Inches(left)
        style.paragraph_format.first_line_indent = Inches(-hanging)
        style.paragraph_format.space_after = Pt(2.5)

    for style_name in ("Caption", "Quote"):
        style = styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    styles["Caption"].font.size = Pt(8.8)
    styles["Caption"].font.italic = True
    styles["Caption"].font.color.rgb = RGBColor.from_string(MUTED)
    styles["Caption"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    styles["Caption"].paragraph_format.keep_with_next = True
    styles["Quote"].font.size = Pt(9.5)
    styles["Quote"].font.color.rgb = RGBColor.from_string(NAVY)
    styles["Quote"].paragraph_format.left_indent = Inches(0.25)
    styles["Quote"].paragraph_format.right_indent = Inches(0.25)
    styles["Quote"].paragraph_format.space_before = Pt(5)
    styles["Quote"].paragraph_format.space_after = Pt(8)

    if "Code Block" not in styles:
        code = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    else:
        code = styles["Code Block"]
    code.font.name = "Menlo"
    code.font.size = Pt(8.1)
    code.font.color.rgb = RGBColor.from_string(NAVY)
    code._element.rPr.rFonts.set(qn("w:eastAsia"), "Menlo")
    code.paragraph_format.left_indent = Inches(0.18)
    code.paragraph_format.right_indent = Inches(0.18)
    code.paragraph_format.space_after = Pt(0)
    code.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE


def configure_section(section, *, first: bool = False) -> None:
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.orientation = WD_ORIENT.PORTRAIT
    section.top_margin = Cm(1.75)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.9)
    section.right_margin = Cm(1.9)
    section.header_distance = Cm(0.65)
    section.footer_distance = Cm(0.65)
    section.different_first_page_header_footer = first


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_run_fonts(run)
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MUTED)
    page = paragraph.add_run()
    add_field(page, "PAGE")
    set_run_fonts(page)
    page.font.size = Pt(8)
    page.font.color.rgb = RGBColor.from_string(MUTED)


def configure_headers(document: Document) -> None:
    section = document.sections[0]
    header = section.header
    table = header.add_table(rows=1, cols=2, width=Inches(6.7))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_fixed(table, [4.55, 2.15])
    set_repeat_table_header(table.rows[0])
    left, right = table.rows[0].cells
    left.text = "HKDSE ICT SCHOOL-BASED ASSESSMENT"
    right.text = "MASTERMIND GAME SYSTEM"
    for cell, alignment in ((left, WD_ALIGN_PARAGRAPH.LEFT), (right, WD_ALIGN_PARAGRAPH.RIGHT)):
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = alignment
        paragraph.paragraph_format.space_after = Pt(0)
        for run in paragraph.runs:
            set_run_fonts(run)
            run.font.size = Pt(7.3)
            run.font.bold = True
            run.font.color.rgb = RGBColor.from_string(TEAL)
    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.text = "CLI-centred technical report  •  Baseline 67b4cc9"
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in paragraph.runs:
        set_run_fonts(run)
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(MUTED)
    tab_stops = paragraph.paragraph_format.tab_stops
    tab_stops.add_tab_stop(Inches(7.0))
    paragraph.add_run("\t")
    add_page_number(paragraph)


def add_cover(document: Document) -> None:
    section = document.sections[0]
    configure_section(section, first=True)
    banner = document.add_table(rows=1, cols=1)
    banner.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_fixed(banner, [6.7])
    set_repeat_table_header(banner.rows[0])
    cell = banner.cell(0, 0)
    set_cell_shading(cell, NAVY)
    set_cell_margins(cell, top=170, bottom=170, start=100, end=100)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("HKDSE ICT  •  SCHOOL-BASED ASSESSMENT")
    set_run_fonts(run)
    run.font.size = Pt(10)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(WHITE)

    document.add_paragraph().paragraph_format.space_after = Pt(26)
    if ICON.exists():
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(18)
        picture = p.add_run().add_picture(str(ICON), width=Inches(1.05))
        picture._inline.docPr.set("descr", "Mastermind Game System application icon")
        picture._inline.docPr.set("title", "Mastermind Game System icon")

    p = document.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    for index, line in enumerate(
        (
            "Hong Kong Diploma of Secondary Education",
            "Information and Communication Technology",
            "School-based Assessment",
        )
    ):
        run = p.add_run(line)
        set_run_fonts(run)
        run.font.size = Pt(17 if index < 2 else 19)
        run.font.color.rgb = RGBColor.from_string(NAVY if index < 2 else TEAL)
        run.font.bold = True
        run.add_break()

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(16)
    title.paragraph_format.space_after = Pt(22)
    run = title.add_run("Mastermind\nGame System")
    set_run_fonts(run)
    run.font.size = Pt(34)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(CORAL)

    rule = document.add_table(rows=1, cols=1)
    rule.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_fixed(rule, [2.0])
    set_repeat_table_header(rule.rows[0])
    set_cell_shading(rule.cell(0, 0), TEAL)
    set_cell_margins(rule.cell(0, 0), top=80, bottom=80, start=100, end=100)
    rule.cell(0, 0).height = Pt(4)

    details = document.add_table(rows=4, cols=2)
    details.alignment = WD_TABLE_ALIGNMENT.CENTER
    details.style = "Table Grid"
    set_table_fixed(details, [2.15, 4.4])
    set_repeat_table_header(details.rows[0])
    values = (
        ("Candidate name", CANDIDATE_NAME),
        ("Candidate number", CANDIDATE_NUMBER),
        ("School", SCHOOL_NAME),
        ("Submission", "1 September 2026  •  CLI-centred system report"),
    )
    for row, pair in zip(details.rows, values, strict=True):
        for index, value in enumerate(pair):
            row.cells[index].text = value
            row.cells[index].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(row.cells[index], top=115, bottom=115)
            if index == 0:
                set_cell_shading(row.cells[index], PALE_BLUE)
            for run in row.cells[index].paragraphs[0].runs:
                set_run_fonts(run)
                run.font.size = Pt(9)
                run.font.bold = index == 0
                run.font.color.rgb = RGBColor.from_string(NAVY if index == 0 else INK)

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    run = p.add_run("DESIGN  •  IMPLEMENTATION  •  TESTING  •  EVALUATION")
    set_run_fonts(run)
    run.font.size = Pt(8.5)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(MUTED)
    p.add_run().add_break(WD_BREAK.PAGE)


def add_front_matter(document: Document) -> None:
    title = document.add_paragraph("Student Declaration and Verification Scope", style="Heading 1")
    title.paragraph_format.page_break_before = False
    add_bookmark(title, "front_declaration", 9000)
    for text in (
        f"This report documents the Mastermind Game System implemented by {CANDIDATE_NAME} ({CANDIDATE_NUMBER}) of {SCHOOL_NAME} in the accompanying repository. The student should confirm authorship, cite all assistance required by school policy, and retain the tested repository baseline.",
        "Technical statements are grounded in source code at commit 67b4cc9be06961483661bf3219626a4ea2e6014e and verification executed on 1 September 2026. This baseline includes the serialized concurrent-duel repair at commit 13fa42f0d9703f945881610a9a2d89ac7d31446f. The report does not claim that every reconstructed debugging scenario corresponds to an individually preserved historical commit; Section 7 identifies the final safeguards and their regression evidence.",
    ):
        p = document.add_paragraph()
        add_inline(p, text)

    table = document.add_table(rows=3, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_fixed(table, [2.25, 4.25])
    set_repeat_table_header(table.rows[0])
    for row, values in zip(
        table.rows,
        (
            ("Student signature", "________________________________________"),
            ("Teacher signature", "________________________________________"),
            ("Date", "________________________________________"),
        ),
        strict=True,
    ):
        for index, value in enumerate(values):
            row.cells[index].text = value
            set_cell_margins(row.cells[index], top=150, bottom=150)
            if index == 0:
                set_cell_shading(row.cells[index], SAND)
            for run in row.cells[index].paragraphs[0].runs:
                set_run_fonts(run)
                run.font.size = Pt(9)
                run.font.bold = index == 0

    heading = document.add_paragraph("Executive Summary", style="Heading 2")
    heading.paragraph_format.space_before = Pt(18)
    summary = "This SBA presents a command-line Mastermind system built in Python 3.13 around a shared canonical rule engine. Players can choose four presets or Custom mode, play against a secure computer Code Maker or a human in pass-and-play, receive duplicate-safe feedback, review local scores, and export validated CSV data. The implementation separates interaction, deterministic rules, and persistence through typed immutable models and dependency injection."
    p = document.add_paragraph()
    add_inline(p, summary)
    p = document.add_paragraph(style="Quote")
    add_inline(
        p,
        "Verification result: **117 Python tests passed** and **36 browser tests passed** with 8 intentional browser skips; focused CLI/core tests achieved **92% measured coverage**, with clean Ruff and mypy checks.",
    )
    document.add_page_break()


def collect_metadata(lines: list[str]) -> tuple[list[tuple[int, str, str]], list[tuple[str, str]]]:
    headings: list[tuple[int, str, str]] = []
    figures: list[tuple[str, str]] = []
    count = 0
    for line in lines:
        match = HEADING_RE.match(line)
        if match:
            count += 1
            headings.append((len(match.group(1)), match.group(2), f"h_{count}"))
        figure = FIGURE_RE.match(line)
        if figure:
            figures.append((figure.group(1), figure.group(2)))
    return headings, figures


def add_toc(
    document: Document, headings: list[tuple[int, str, str]], page_map: dict[str, int]
) -> None:
    title = document.add_paragraph("Table of Contents", style="Heading 1")
    title.paragraph_format.page_break_before = False
    intro = document.add_paragraph(
        "The entries below are linked to their sections. Page numbers are refreshed from the rendered report during the build process."
    )
    intro.style = "Quote"
    for level, text, anchor in headings:
        p = document.add_paragraph()
        p.paragraph_format.left_indent = Inches((level - 1) * 0.24)
        p.paragraph_format.space_after = Pt(1.1)
        p.paragraph_format.keep_together = True
        p.paragraph_format.tab_stops.add_tab_stop(Inches(6.85), 2, 1)
        add_internal_link(p, text, anchor, NAVY if level == 1 else INK)
        if p.runs:
            for run in p.runs:
                run.font.size = Pt(9 if level == 1 else 8.3)
                run.font.bold = level == 1
        p.add_run("\t")
        page_run = p.add_run(str(page_map.get(anchor, "—")))
        set_run_fonts(page_run)
        page_run.font.size = Pt(8.3)
        page_run.font.bold = level == 1
        page_run.font.color.rgb = RGBColor.from_string(NAVY)
    document.add_page_break()


def add_lists_page(document: Document, figures: list[tuple[str, str]]) -> None:
    title = document.add_paragraph("List of Figures and Tables", style="Heading 1")
    title.paragraph_format.page_break_before = False
    document.add_paragraph("Figures", style="Heading 2")
    for _, caption in figures:
        p = document.add_paragraph(style="List Bullet")
        add_inline(p, caption)
    document.add_paragraph("Principal tables", style="Heading 2")
    for text in (
        "Minimum expectations and delivered evidence",
        "Language features and system applications",
        "General parameters and preset configuration",
        "Core data structures, project configuration, and libraries",
        "Testing plan, cases, constraints, and acceptance results",
        "Debugging error list and improvement priorities",
        "Reference register and Gantt-phase completion evidence",
    ):
        p = document.add_paragraph(style="List Bullet")
        add_inline(p, text)
    document.add_page_break()


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        line = lines[index].strip()
        if not TABLE_SEPARATOR_RE.match(line):
            rows.append([cell.strip() for cell in line.strip("|").split("|")])
        index += 1
    return rows, index


def add_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    columns = max(len(row) for row in rows)
    normalized = [row + [""] * (columns - len(row)) for row in rows]
    table = document.add_table(rows=len(normalized), cols=columns)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    total = 6.7
    widths = [total / columns] * columns
    if columns == 2:
        widths = [2.15, 4.55]
    elif columns == 3:
        widths = [1.6, 2.75, 2.35]
    elif columns == 4:
        widths = [1.35, 2.15, 1.3, 1.9]
    elif columns == 5:
        widths = [0.95, 1.55, 1.3, 1.3, 1.6]
    set_table_fixed(table, widths)
    for row_index, row in enumerate(normalized):
        if row_index == 0:
            set_repeat_header(table.rows[row_index])
        else:
            tr_pr = table.rows[row_index]._tr.get_or_add_trPr()
            tr_pr.append(OxmlElement("w:cantSplit"))
        for column_index, value in enumerate(row):
            cell = table.cell(row_index, column_index)
            cell.text = ""
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_index == 0:
                set_cell_shading(cell, NAVY)
            elif row_index % 2 == 0:
                set_cell_shading(cell, PALE_BLUE)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_inline(p, value)
            for run in p.runs:
                set_run_fonts(run)
                run.font.size = Pt(8.15)
                if row_index == 0:
                    run.font.bold = True
                    run.font.color.rgb = RGBColor.from_string(WHITE)
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def add_figure(document: Document, filename: str, caption: str) -> None:
    path = ASSETS / filename
    if not path.exists():
        raise FileNotFoundError(path)
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    picture = p.add_run().add_picture(str(path), width=Inches(6.65))
    picture._inline.docPr.set("descr", caption)
    picture._inline.docPr.set("title", caption.split(".", 1)[0])
    cap = document.add_paragraph(style="Caption")
    cap.add_run(caption)


def set_paragraph_shading(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


def add_body(document: Document, lines: list[str]) -> None:
    bookmark_id = 1
    heading_number = 0
    ordered_num_id: int | None = None
    index = 0
    in_code = False
    while index < len(lines):
        raw = lines[index]
        line = raw.rstrip()
        if line.startswith("```"):
            in_code = not in_code
            ordered_num_id = None
            index += 1
            continue
        if in_code:
            ordered_num_id = None
            p = document.add_paragraph(style="Code Block")
            p.paragraph_format.keep_together = True
            set_paragraph_shading(p, "F1F5F7")
            p.add_run(line or " ")
            index += 1
            continue
        if not line.strip():
            ordered_num_id = None
            index += 1
            continue
        figure = FIGURE_RE.match(line)
        if figure:
            ordered_num_id = None
            add_figure(document, figure.group(1), figure.group(2))
            index += 1
            continue
        heading = HEADING_RE.match(line)
        if heading:
            ordered_num_id = None
            level = len(heading.group(1))
            heading_number += 1
            p = document.add_paragraph(heading.group(2), style=f"Heading {level}")
            add_bookmark(p, f"h_{heading_number}", bookmark_id)
            bookmark_id += 1
            index += 1
            continue
        if (
            line.strip().startswith("|")
            and index + 1 < len(lines)
            and TABLE_SEPARATOR_RE.match(lines[index + 1].strip())
        ):
            ordered_num_id = None
            rows, index = parse_table(lines, index)
            add_table(document, rows)
            continue
        if line.startswith("> "):
            ordered_num_id = None
            p = document.add_paragraph(style="Quote")
            add_inline(p, line[2:])
            index += 1
            continue
        bullet = re.match(r"^-\s+(.+)$", line)
        if bullet:
            ordered_num_id = None
            p = document.add_paragraph(style="List Bullet")
            add_inline(p, bullet.group(1))
            index += 1
            continue
        number = re.match(r"^\d+\.\s+(.+)$", line)
        if number:
            if ordered_num_id is None:
                ordered_num_id = create_numbering_instance(document)
            p = document.add_paragraph(style="List Number")
            apply_numbering(p, ordered_num_id)
            add_inline(p, number.group(1))
            index += 1
            continue
        ordered_num_id = None
        p = document.add_paragraph()
        add_inline(p, line)
        index += 1


def build(output: Path, page_map_path: Path | None = None) -> None:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    headings, figures = collect_metadata(lines)
    page_map = {}
    if page_map_path and page_map_path.exists():
        page_map = json.loads(page_map_path.read_text(encoding="utf-8"))

    document = Document()
    configure_styles(document)
    add_cover(document)
    configure_headers(document)
    add_front_matter(document)
    add_toc(document, headings, page_map)
    add_lists_page(document, figures)
    add_body(document, lines)

    props = document.core_properties
    props.title = (
        "Hong Kong Diploma of Secondary Education — Information and Communication Technology — "
        "School-based Assessment — Mastermind Game System"
    )
    props.subject = "CLI-centred Mastermind system analysis, design, implementation, and evaluation"
    props.author = CANDIDATE_NAME
    props.keywords = "HKDSE, ICT, SBA, Mastermind, CLI, Python, testing"
    props.comments = "Generated reproducibly from report_source.md and repository evidence."
    props.last_modified_by = CANDIDATE_NAME
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--page-map", type=Path)
    args = parser.parse_args()
    build(args.output.resolve(), args.page_map.resolve() if args.page_map else None)


if __name__ == "__main__":
    main()
